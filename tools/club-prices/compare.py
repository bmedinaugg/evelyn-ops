#!/usr/bin/env python3
"""Every club's prices, four ways: the form, the bot's copy of it, Magicline, join-now.

    python3 tools/club-prices/compare.py docs/club-pricing-comparison.xlsx

WHAT THE THREE SOURCES ARE
  1. The CHANGE FORM  — Freshdesk cf_clubs, read live. Source of truth by
     decision (2026-09-21): where it disagrees with Magicline, the form wins.
  1b. WHAT THE BOT QUOTES — bot.form_schemas in Supabase. The bot does NOT read
     Freshdesk; its "Fetch Form Options" node reads this MIRROR. So a row where
     the mirror disagrees with Freshdesk is a straight sync bug: the form is
     right and the member is told something else. Found on the first run
     (Den Haag Dagelijkse Groenmarkt HOME+, mirror EUR5 low on three terms),
     which is why the mirror is a column and not an assumption.
  2. MAGICLINE        — GET /connect/v1/rate-bundle?studioId=. Every rate that
     exists as a sellable product at that club.
  3. JOIN-NOW         — trainmore.com/en-NL/join-now. What a NEW member is
     actually offered.

WHY THERE IS NO JOIN-NOW COLUMN, ONLY A LINK
join-now reads the SAME rate-bundle endpoint (confirmed live 2026-09-21 by
watching the network: GET /connect/v2/studio, GET /connect/v2/studio/{id},
GET /connect/v1/rate-bundle?studioId= — no auth, and no /preview call while
browsing; the fake-customer preview POST only happens at checkout for
age-adjusted or voucher pricing). But the site then FILTERS that list with a
rule of its own -- the page says "Access label is based on the label of your
chosen home club" -- and that rule is not in the API. Verified by hand at
Den Haag Laan van NOI: Magicline has CITY+, HOME+ and PREMIUM; join-now offers
CITY+ and HOME+ only. Nothing on the bundle objects (limitedOfferingPeriod,
preuseType, rateCodeDto, contractStartType...) separates the shown from the
hidden, so a computed column would be a guess dressed as data.

So each row carries the join-now URL for that exact club and tier. Click it and
the page either offers that tier or it does not -- one click, no inference.

THE COMPARISON IS NOT AN ERROR LIST
A change-form row differing from Magicline is not automatically wrong. New joins
and membership changes are different products and may legitimately differ. This
sheet says where they disagree; only Member Care can say which are intended.

Student and corporate bundles are excluded throughout: neither is sold via this
form. Reads only -- nothing here writes anywhere.
"""
import base64, json, os, re, sys, time, urllib.parse, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/club-pricing-comparison.xlsx"
ML = "https://ugg.api.magicline.com/connect"
JOIN = "https://trainmore.com/en-NL/join-now/"

for line in open(".env.local"):
    line = line.strip()
    if "=" in line and not line.startswith("#"):
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
AUTH = base64.b64encode(f"{os.environ['FRESHDESK_API_KEY']}:X".encode()).decode()


def fd(path):
    req = urllib.request.Request(f"https://urbangymgroup.freshdesk.com{path}",
                                 headers={"Authorization": f"Basic {AUTH}"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def ml(path):
    req = urllib.request.Request(ML + path, headers={"Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


LABEL = re.compile(r"\((Red|Black|Regular) Label\)", re.I)
OPENS = re.compile(r"[-–]\s*opens.*$", re.I)
SKIP_BUNDLE = re.compile(r"corporate|pay in full|\bSTU\b|intermediar|invoice|split", re.I)
TERMS = {(1, "YEAR"): "1 year", (2, "YEAR"): "2 years",
         (3, "YEAR"): "3 years", (4, "WEEK"): "Flex"}
TERM_ORDER = {"1 year": 0, "2 years": 1, "3 years": 2, "Flex": 3}
TIER_ORDER = {"HOME": 0, "HOME+": 1, "CITY+": 2, "PREMIUM": 3}
PRICE = re.compile(r"€\s*([0-9]+(?:[.,][0-9]{1,2})?)")
YEARS = re.compile(r"(\d+)\s*[-–—]?\s*year", re.I)


def norm_club(name):
    s = OPENS.sub("", LABEL.sub("", name))
    s = re.sub(r"\btrainmore\b", "", s, flags=re.I)
    s = re.sub(r"\bb\.?v\.?\b", "", s, flags=re.I)
    s = re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()
    for a, b in (("van woustraat", "van wou"), ("oost blacklabel", "oost"),
                 ("de savornin", "savornin"), ("west ladies", "west"),
                 ("bussum landstraat", "bussum"), ("bilthoven leyenseweg", "bilthoven"),
                 ("leiden breestraat", "leiden"), ("leusden maximaplein", "leusden")):
        if s == a:
            s = b
    return s


def norm_tier(t):
    u = t.upper().replace(" ", "")
    if "CITY+" in u: return "CITY+"
    if "HOME+" in u: return "HOME+"
    if "PREMIUM" in u: return "PREMIUM"
    if re.search(r"\bHOME\b", t, re.I) or u.startswith("HOME"): return "HOME"
    return None


def parse_option(text):
    m = PRICE.search(text)
    price = float(m.group(1).replace(",", ".")) if m else None
    if re.search(r"flex", text, re.I):
        return "Flex", price
    y = YEARS.search(text)
    return (f"{y.group(1)} year" + ("s" if y.group(1) != "1" else ""), price) if y else (None, price)


def join_url(studio_id, bundle_id, tier):
    return JOIN + "?" + urllib.parse.urlencode({
        "studioId": studio_id, "type": "Regular",
        "rateBundleId": bundle_id, "accessLevelName": tier.lower(), "workouts": 0})


# ---- read both readable sources ------------------------------------------
def supabase_mirror():
    """cf_clubs as the BOT sees it. Failure here must not lose the whole run."""
    try:
        url, key = os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_ROLE_KEY"]
        req = urllib.request.Request(
            f"{url}/rest/v1/form_schemas?field_key=eq.cf_clubs&select=options",
            headers={"apikey": key, "Authorization": f"Bearer {key}", "Accept-Profile": "bot"})
        with urllib.request.urlopen(req, timeout=60) as r:
            rows = json.load(r)
        return ((rows[0].get("options") or {}).get("choices") or {}) if rows else {}
    except Exception as e:
        print(f"  mirror unreadable ({e}) — that column will be blank", file=sys.stderr)
        return {}


print("reading Freshdesk…", file=sys.stderr)
tree = {f["name"]: f for f in fd("/api/v2/ticket_fields")}["cf_clubs"]["choices"]
form = {}                                   # club -> tier -> term -> (price, raw, display)
for club, tiers in tree.items():
    key = norm_club(club)
    for tier, options in (tiers or {}).items():
        t = norm_tier(tier)
        if not t:
            continue
        for opt in options or []:
            term, price = parse_option(opt)
            if term:
                form.setdefault(key, {}).setdefault(t, {})[term] = (price, opt, club)

print("reading the bot's mirror…", file=sys.stderr)
botq = {}                                   # club -> tier -> term -> price
for club, tiers in supabase_mirror().items():
    key = norm_club(club)
    for tier, options in (tiers or {}).items():
        t = norm_tier(tier)
        if not t:
            continue
        for opt in options or []:
            term, price = parse_option(opt)
            if term:
                botq.setdefault(key, {}).setdefault(t, {})[term] = price

print("reading Magicline…", file=sys.stderr)
studios = [s for s in ml("/v2/studio")
           if s["studioName"].lower().startswith("trainmore")
           and s["address"]["countryCodeAlpha2"] == "NL"
           and s["studioName"] != "TrainMore HQ"]
mag, sid_of, bundle_of = {}, {}, {}
for s in studios:
    key = norm_club(s["studioName"])
    sid_of[key] = s["id"]
    try:
        bundles = ml(f"/v1/rate-bundle?studioId={s['id']}")
    except Exception:
        continue
    for b in bundles:
        if SKIP_BUNDLE.search(b["name"]):
            continue
        t = norm_tier(b["name"])
        if not t:
            continue
        bundle_of[(key, t)] = b["id"]
        for term in b["terms"]:
            label = TERMS.get((term["termValue"], term["termUnit"]))
            if label:
                mag.setdefault(key, {}).setdefault(t, {})[label] = term["price"]
    time.sleep(0.12)

# ---- build one row per club / tier / term ---------------------------------
rows = []
for club in sorted(set(form) | set(mag)):
    f, m = form.get(club, {}), mag.get(club, {})
    display = next((v[2] for t in f.values() for v in t.values()), None) \
        or next((s["studioName"] for s in studios if norm_club(s["studioName"]) == club), club.title())
    sid = sid_of.get(club, "")
    bq = botq.get(club, {})
    for tier in sorted(set(f) | set(m), key=lambda t: TIER_ORDER.get(t, 9)):
        ft, mt, bt = f.get(tier, {}), m.get(tier, {}), bq.get(tier, {})
        link = join_url(sid, bundle_of[(club, tier)], tier) if sid and (club, tier) in bundle_of else ""
        for term in sorted(set(ft) | set(mt), key=lambda t: TERM_ORDER.get(t, 9)):
            fv = ft.get(term)
            mv = mt.get(term)
            bv = bt.get(term)
            fp = fv[0] if fv else None
            # The mirror is checked FIRST and outranks everything else: the other
            # statuses are questions for Member Care, this one is just broken.
            stale = (botq and fp is not None
                     and (bv is None or abs(bv - fp) > 0.005))
            if stale:
                status = "BOT OUT OF SYNC"
                action = (f'Bot quotes €{bv:g}, the form says €{fp:g} — resync bot.form_schemas'
                          if bv is not None else
                          'The form has this row and the bot\'s copy does not — resync bot.form_schemas')
            elif fv and mv is not None and fp is not None and abs(fp - mv) > 0.005:
                status, action = "price differs", f'Form says €{fp:g}, Magicline €{mv:g}'
            elif fv and mv is None:
                status, action = "form only", "On the form, not sold in Magicline — check join-now"
            elif mv is not None and not fv:
                status, action = "MISSING from the form", f'Magicline sells this at €{mv:g} — should the form offer it?'
            else:
                status, action = "match", ""
            rows.append([display, tier, term,
                         (f"€{fp:g}" if fp is not None else ("—" if fv else "")),
                         (f"€{bv:g}" if bv is not None else ""),
                         (f"€{mv:g}" if mv is not None else ""),
                         status, action, link])

ORDER = {"BOT OUT OF SYNC": 0, "MISSING from the form": 1, "price differs": 2, "form only": 3}
diffs = sorted([r for r in rows if r[6] != "match"],
               key=lambda r: (ORDER.get(r[6], 9), r[0], TIER_ORDER.get(r[1], 9)))

# ---- write ----------------------------------------------------------------
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F3E4E")
FILLS = {"price differs": PatternFill("solid", fgColor="FDE9D9"),
         "MISSING from the form": PatternFill("solid", fgColor="FCE4E4"),
         "form only": PatternFill("solid", fgColor="FFF7D6"),
         "BOT OUT OF SYNC": PatternFill("solid", fgColor="F4C7C3")}
LINK = Font(color="0563C1", underline="single")

wb = Workbook()
ws = wb.active
ws.title = "Read me"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 104
for a, b in [
    ("Club pricing: three sources", ""),
    ("", ""),
    ("1. The change form", "Freshdesk cf_clubs — what the bot offers an EXISTING member changing membership. "
                           "Source of truth by decision (21 Sep 2026): where it disagrees with Magicline, the form wins."),
    ("1b. What the bot quotes", "bot.form_schemas in Supabase. The bot does NOT read Freshdesk — its "
                                "\"Fetch Form Options\" node reads this mirror. Where this column disagrees with "
                                "the form, the member is being told the wrong price and nobody has to decide "
                                "anything: it is a sync bug. Those rows sort to the top of Differences."),
    ("2. Magicline", "GET /connect/v1/rate-bundle — every rate that exists as a sellable product at that club. "
                     "Student and corporate bundles are excluded; neither is sold through this form."),
    ("3. join-now", "trainmore.com/en-NL/join-now — what a NEW member is actually offered. "
                    "There is no join-now COLUMN because the site reads the same Magicline endpoint and then "
                    "filters it with a rule of its own (\"access label is based on the label of your chosen home "
                    "club\") that is not exposed in the API. A computed column would be a guess. Instead every "
                    "row carries a link: click it and the page either offers that tier or it does not."),
    ("", ""),
    ("This is not an error list", "A difference is not automatically a mistake. New joins and membership changes "
                                  "are different products and may legitimately differ. This sheet says where they "
                                  "disagree; only Member Care can say which differences are intended."),
    ("How to use it", "Work the Differences tab. Each row names the club, tier and term, both prices, and what to "
                      "check. The link opens join-now for that exact club and tier."),
    ("Billing frequency", "Every consumer rate in Magicline bills every 4 weeks, all four contract lengths. "
                          "The form labels do not say so, which is why the bot used to guess."),
    ("", ""),
    ("Rows compared", str(len(rows))),
    ("Differences", str(len(diffs))),
]:
    ws.append([a, b])
ws["A1"].font = Font(bold=True, size=15)
for r in range(3, ws.max_row + 1):
    ws.cell(r, 1).font = Font(bold=True)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if r not in (5, 7) else 62

HEADERS = ["Club", "Tier", "Term", "Change form (Freshdesk)", "What the bot quotes",
           "Magicline", "Status", "What to check", "Open join-now for this club + tier"]
for title, data in (("Differences", diffs), ("All rows", rows)):
    ws = wb.create_sheet(title)
    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in data:
        ws.append(r[:8] + [""])
        cell = ws.cell(ws.max_row, 9)
        if r[8]:
            cell.value = "check on join-now"
            cell.hyperlink = r[8]
            cell.font = LINK
    for i, w in enumerate([38, 9, 9, 17, 16, 12, 22, 54, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    for i, r in enumerate(data, start=2):
        fill = FILLS.get(r[6])
        if fill:
            for c in range(1, 9):
                ws.cell(i, c).fill = fill
        ws.cell(i, 8).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
by_status = {}
for r in rows:
    by_status[r[6]] = by_status.get(r[6], 0) + 1
print(OUT)
print(f"{len(rows)} rows across {len({r[0] for r in rows})} clubs | " +
      " | ".join(f"{k}: {v}" for k, v in sorted(by_status.items())))
