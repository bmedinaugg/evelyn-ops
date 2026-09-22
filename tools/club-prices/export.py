#!/usr/bin/env python3
"""Export every club price Freshdesk holds into an Excel workbook.

Prices are not stored anywhere as data — they are typed into the *labels* of
Freshdesk dropdown options, e.g. "1 - year €92". This reads those fields live
from the Freshdesk API and turns them into a spreadsheet.

    python3 tools/club-prices/export.py [output.xlsx]

Needs FRESHDESK_API_KEY in .env.local, and openpyxl (pip3 install openpyxl).
"""
import base64, json, os, re, sys, urllib.request
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DOMAIN = "urbangymgroup.freshdesk.com"
OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/club-prices.xlsx"

# The three nested dropdowns that carry prices, and the club list to check them
# against. Verified against /api/v2/ticket_fields: these are the only fields
# whose choices contain a euro sign.
PRICE_FIELDS = [
    ("cf_clubs", "TrainMore", "Membership change"),
    ("cf_club_where_they_want_to_extend_at", "TrainMore", "Extension"),
    ("cf_access_level_extension_request_clubsportive", "Clubsportive", "Extension"),
]
MASTER_FIELD = "cf_club"


# ------------------------------------------------------------------ fetch
def load_env(path=".env.local"):
    if os.path.exists(path):
        for line in open(path):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def ticket_fields():
    key = os.environ.get("FRESHDESK_API_KEY")
    if not key:
        sys.exit("FRESHDESK_API_KEY not set (expected in .env.local)")
    auth = base64.b64encode(f"{key}:X".encode()).decode()
    req = urllib.request.Request(
        f"https://{DOMAIN}/api/v2/ticket_fields",
        headers={"Authorization": f"Basic {auth}"},
    )
    with urllib.request.urlopen(req, timeout=60) as r:
        fields = json.load(r)
    if not isinstance(fields, list) or not fields:
        sys.exit("Freshdesk returned no ticket fields")
    return {f["name"]: f for f in fields}


# ------------------------------------------------------------ parse labels
PRICE = re.compile(r"€\s*([0-9]+(?:[.,][0-9]{1,2})?)")
PRICE_TRAILING = re.compile(r"([0-9]+(?:[.,][0-9]{1,2})?)\s*€")
BARE = re.compile(r"years?\s*[-–—:]?\s*([0-9]+(?:[.,][0-9]{1,2})?)\s*$", re.I)
YEARS = re.compile(r"(\d+)\s*[-–—]?\s*year", re.I)
USUALLY = re.compile(r"usually[:\s]*€\s*([0-9]+(?:[.,][0-9]{1,2})?)", re.I)
FOUR_WEEKS = re.compile(r"per\s*4\s*weeks", re.I)
LABEL = re.compile(r"\((Red|Black|Regular) Label\)", re.I)
OPENS = re.compile(r"[-–]\s*opens.*$", re.I)


def money(s):
    return float(s.replace(",", "."))


def parse_option(text):
    """Pull structured price data out of one dropdown option label."""
    issue = ""
    amounts = PRICE.findall(text)
    if amounts:
        price = money(amounts[0])
    elif PRICE_TRAILING.search(text):
        price, issue = money(PRICE_TRAILING.search(text).group(1)), "€ sign written after the amount"
    elif BARE.search(text):
        price, issue = money(BARE.search(text).group(1)), "no € sign at all"
    else:
        price, issue = None, "no price in the option text"

    if re.search(r"flex", text, re.I):
        duration = "Flex"
    else:
        m = YEARS.search(text)
        duration = f"{m.group(1)} year" + ("s" if m.group(1) != "1" else "") if m else None
    if duration is None:
        issue = issue or "duration not recognisable"

    was = USUALLY.search(text)
    standard = money(was.group(1)) if was else None
    period = "per 4 weeks" if FOUR_WEEKS.search(text) else "per month"

    note = "; ".join(re.findall(r"\(([^)]*)\)", text))
    if not issue and note and not was:
        issue = "extra amount or condition in the text"
    elif not issue and len(amounts) > 1 and not was:
        issue = "more than one amount in the text"
    return duration, price, standard, period, note, issue


def norm(club):
    """Normalise a club name so the different fields can be matched up."""
    s = re.sub(r"[^a-z0-9]", "", OPENS.sub("", LABEL.sub("", club)).lower())
    for variant, canonical in (("vanwoustraat", "vanwou"), ("oostblacklabel", "oost"),
                               ("desavornin", "savornin"), ("westladies", "west"),
                               ("bussumlandstraat", "bussum"),
                               ("bilthovenleyenseweg", "bilthoven")):
        s = s.replace(variant, canonical)
    return s


def flatten(field, brand, form):
    """nested_field choices are {club: {access level: [option, ...]}}."""
    choices = field.get("choices")
    if not isinstance(choices, dict):
        sys.exit(f"{field['name']}: expected a nested field, got {type(choices).__name__}")
    rows = []
    for club in sorted(choices):
        for access, options in (choices[club] or {}).items():
            for text in options or []:
                duration, price, standard, period, note, issue = parse_option(text)
                m = LABEL.search(club)
                rows.append({
                    "brand": brand, "form": form,
                    "city": club.split()[0], "club": club,
                    "label": m.group(1).title() + " Label" if m else "",
                    "access": access.split(" (")[0].replace("CITY +", "CITY+").strip(),
                    "access_full": access,
                    "duration": duration, "price": price, "standard": standard,
                    "period": period, "note": note, "raw": text, "issue": issue,
                })
    return rows


# ---------------------------------------------------------------- workbook
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F3E4E")
WARN = PatternFill("solid", fgColor="FDE9D9")
EUR = '€#,##0.00'


def build(rows, master, fields, fetched):
    wb = Workbook()

    def sheet(title, headers, data, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = HEAD, HEAD_FILL
        for r in data:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    # -- Read me ----------------------------------------------------------
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 98
    src = "; ".join(f"{n} (id {fields[n]['id']})" for n, _, _ in PRICE_FIELDS)
    for a, b in [
        ("Club prices", ""),
        ("", ""),
        ("Source", f"Freshdesk API — https://{DOMAIN}/api/v2/ticket_fields"),
        ("Fields read", src),
        ("Fetched", fetched),
        ("", ""),
        ("How prices are stored", "Not as data. Each price is typed into the TEXT of a dropdown "
                                  "option, e.g. \"1 - year €92\". The 'Raw Freshdesk option' column "
                                  "is that text exactly — it is also what the bot quotes to members."),
        ("To change a price", "Edit the dropdown option in Freshdesk (Admin → Ticket Fields). The "
                              "bot picks it up on the next nightly sync (04:00 → Supabase "
                              "ticket_taxonomy, 06:00 → form_schemas)."),
        ("Shaded rows", "Option text written inconsistently — see the 'Data quality' sheet."),
        ("", ""),
        ("Sheets", "All prices — every priced option from all three fields"),
        ("", "Change matrix — TrainMore change-form prices as a club × access level × duration grid"),
        ("", "Data quality — option text that does not follow the usual pattern"),
        ("", "Coverage — every club on the cf_club master list vs. what it has priced"),
    ]:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=15)
    for r in range(3, ws.max_row + 1):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32

    # -- All prices -------------------------------------------------------
    ws = sheet("All prices",
               ["Brand", "Form", "City", "Club", "Club label", "Access level", "Duration",
                "Price", "Per", "Standard price", "Note", "Raw Freshdesk option"],
               [[r["brand"], r["form"], r["city"], r["club"], r["label"], r["access"],
                 r["duration"], r["price"], r["period"], r["standard"], r["note"], r["raw"]]
                for r in rows],
               [12, 18, 12, 44, 14, 12, 10, 11, 12, 14, 22, 40])
    for i, r in enumerate(rows, start=2):
        ws.cell(i, 8).number_format = EUR
        ws.cell(i, 10).number_format = EUR
        if r["issue"]:
            for c in range(1, 13):
                ws.cell(i, c).fill = WARN

    # -- Change matrix ----------------------------------------------------
    change = [r for r in rows if r["form"] == "Membership change"]
    order = ["HOME", "HOME+", "CITY+", "PREMIUM"]
    accesses = sorted({r["access"] for r in change},
                      key=lambda a: (order.index(a) if a in order else 99, a))
    durations = ["1 year", "2 years", "3 years", "Flex"]
    cols = [(a, d) for a in accesses for d in durations]
    lookup = {(r["club"], r["access"], r["duration"]): r["price"] for r in change}

    ws = wb.create_sheet("Change matrix")
    ws.append([""] + [a for a, _ in cols])
    ws.append(["Club"] + [d for _, d in cols])
    for club in sorted({r["club"] for r in change}):
        ws.append([club] + [lookup.get((club, a, d)) for a, d in cols])
    for c in ws[1] + ws[2]:
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    start = 2
    for i in range(2, len(cols) + 3):
        if i == len(cols) + 2 or ws.cell(1, i).value != ws.cell(1, start).value:
            if i - start > 1:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i - 1)
            start = i
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 38
    for i in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 10
    for row in ws.iter_rows(min_row=3, min_col=2):
        for c in row:
            c.number_format = EUR

    # -- Data quality -----------------------------------------------------
    sheet("Data quality",
          ["Form", "Club", "Access level", "Raw Freshdesk option", "Issue", "Price read as"],
          [[r["form"], r["club"], r["access"], r["raw"], r["issue"],
            r["price"] if r["price"] is not None else ""] for r in rows if r["issue"]],
          [20, 46, 14, 38, 34, 14])

    # -- Coverage ---------------------------------------------------------
    by_norm = {}
    for r in rows:
        if r["brand"] != "TrainMore":
            continue
        e = by_norm.setdefault(norm(r["club"]), {"change": set(), "ext": set(), "names": set()})
        e["names"].add(r["club"])
        e["change" if r["form"] == "Membership change" else "ext"].add(r["access"])

    cov, seen = [], set()
    for club in master:
        n = norm(club)
        seen.add(n)
        e = by_norm.get(n)
        ch = sorted(e["change"]) if e else []
        ex = sorted(e["ext"]) if e else []
        names = sorted(e["names"]) if e else []
        if not e:
            verdict = "NO PRICES ANYWHERE — the bot has nothing to quote for this club"
        elif not ch:
            verdict = "missing from the membership-change form"
        elif not ex:
            verdict = "missing from the extension form"
        elif len(names) > 1:
            verdict = "named differently in the two forms: " + " / ".join(names)
        else:
            verdict = ""
        cov.append([club, ", ".join(ch) or "—", len(ch), ", ".join(ex) or "—", len(ex), verdict])
    for n, e in sorted(by_norm.items()):
        if n not in seen:
            cov.append([" / ".join(sorted(e["names"])), ", ".join(sorted(e["change"])) or "—",
                        len(e["change"]), ", ".join(sorted(e["ext"])) or "—", len(e["ext"]),
                        "priced, but not on the cf_club master list"])

    ws = sheet("Coverage",
               ["Club (cf_club master list)", "Access levels — change form", "#",
                "Access levels — extension form", "#", "What to check"],
               cov, [40, 30, 5, 30, 5, 56])
    for i, row in enumerate(cov, start=2):
        if row[5]:
            for c in range(1, 7):
                ws.cell(i, c).fill = WARN

    return wb


def main():
    load_env()
    fields = ticket_fields()
    fetched = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    rows = []
    for name, brand, form in PRICE_FIELDS:
        if name not in fields:
            sys.exit(f"Freshdesk no longer has a field called {name}")
        rows += flatten(fields[name], brand, form)
    if not rows:
        sys.exit("No priced options found — refusing to write an empty workbook")

    master = [c for c in fields[MASTER_FIELD].get("choices", []) if " - " in c]

    os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
    build(rows, master, fields, fetched).save(OUT)

    flagged = sum(1 for r in rows if r["issue"])
    clubs = len({r["club"] for r in rows})
    print(f"{OUT}\n{len(rows)} prices across {clubs} clubs | {flagged} flagged | fetched {fetched}")


if __name__ == "__main__":
    main()
