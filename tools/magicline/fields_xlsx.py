#!/usr/bin/env python3
"""Build the Magicline field inventory workbook.

    python3 tools/magicline/connect_fields.py /tmp/connect_fields.json
    python3 tools/magicline/fields_xlsx.py /tmp/connect_fields.json docs/magicline-fields.xlsx

Connect API fields come from a live sweep; Open API fields come from
tools/magicline/openapi_fields.json, captured from real n8n executions with
personal values redacted.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CONNECT = sys.argv[1] if len(sys.argv) > 1 else "docs/magicline-connect-fields.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else "docs/magicline-fields.xlsx"
OPENAPI = os.path.join(os.path.dirname(os.path.abspath(__file__)), "openapi_fields.json")

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F3E4E")
WARN = PatternFill("solid", fgColor="FDE9D9")
PII = PatternFill("solid", fgColor="EEE6F5")

connect = json.load(open(CONNECT))
openapi = json.load(open(OPENAPI))
wb = Workbook()


def sheet(title, headers, rows, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in wrap_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, i).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def short(v, n=70):
    if v is None:
        return ""
    s = json.dumps(v, ensure_ascii=False) if not isinstance(v, str) else v
    return s if len(s) <= n else s[:n - 1] + "…"


# ---------------------------------------------------------------- Read me
ws = wb.active
ws.title = "Read me"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 104
n_connect = sum(len(e["fields"]) for e in connect["endpoints"])
n_open = sum(len(e["fields"]) for e in openapi["endpoints"])
for a, b in [
    ("Magicline field inventory", ""),
    ("", ""),
    ("What this is", "Every field Magicline actually returned, with a real example value and what "
                     "a member could ask that the field answers."),
    ("Connect API", f"{n_connect} field paths across {len(connect['endpoints'])} endpoints. No key "
                    f"needed. Swept live by tools/magicline/connect_fields.py against "
                    f"{connect['sample_studio']['name']}."),
    ("Open API", f"{n_open} fields across {len(openapi['endpoints'])} endpoints. Needs a per-studio "
                 "x-api-key. Captured from real n8n executions."),
    ("", ""),
    ("Personal data", "Open API examples for anything identifying a person are REDACTED and shown "
                      "as «placeholder». Those rows are shaded. The real values were read once to "
                      "learn the shape and are not reproduced here or in the repo."),
    ("Always null", "Connect fields that came back null on every record are flagged. They exist in "
                    "the schema but this tenant does not populate them — do not build on them "
                    "without checking another club."),
    ("", ""),
    ("Sheets", "Questions — what can be answered, and the field that answers it"),
    ("", "Connect fields — club, product and price fields (no login)"),
    ("", "Open API fields — member fields (login required)"),
    ("", "Connect endpoints — what was called, and what came back"),
]:
    ws.append([a, b])
ws["A1"].font = Font(bold=True, size=15)
for r in range(3, ws.max_row + 1):
    ws.cell(r, 1).font = Font(bold=True)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30

# ------------------------------------------------------------- Questions
questions = []
for e in openapi["endpoints"]:
    for f in e["fields"]:
        if f.get("answers") and "?" in f["answers"]:
            questions.append(["Member (login)", f["answers"], e["endpoint"], f["field"],
                              "" if f.get("personal") else short(f.get("example"), 40)])
for e in connect["endpoints"]:
    if e["status"] == "ok" and e["fields"]:
        questions.append(["Club / product (no login)", e["answers"], e["endpoint"],
                          f"{len(e['fields'])} fields", e["path"]])
ws = sheet("Questions",
           ["Who is asking", "What it answers", "Endpoint", "Field", "Example / path"],
           questions, [24, 62, 24, 40, 44], wrap_cols=(2,))

# --------------------------------------------------------- Connect fields
rows = []
for e in connect["endpoints"]:
    for f in e["fields"]:
        rows.append([e["endpoint"], e["path"], f["field"], f["type"],
                     short(f.get("example")), "yes" if f["always_null"] else ""])
ws = sheet("Connect fields",
           ["Endpoint", "Path", "Field", "Type", "Example value", "Always null"],
           rows, [26, 46, 46, 16, 56, 12])
for i, r in enumerate(rows, start=2):
    if r[5]:
        for c in range(1, 7):
            ws.cell(i, c).fill = WARN

# --------------------------------------------------------- Open API fields
rows = []
for e in openapi["endpoints"]:
    for f in e["fields"]:
        rows.append([e["endpoint"], e["path"], e["tool"], f["field"], f["type"],
                     short(f.get("example")), "yes" if f.get("personal") else "",
                     f.get("answers", "")])
ws = sheet("Open API fields",
           ["Endpoint", "Path", "n8n tool", "Field", "Type", "Example value",
            "Personal", "What it answers"],
           rows, [22, 44, 26, 52, 18, 40, 10, 62], wrap_cols=(8,))
for i, r in enumerate(rows, start=2):
    if r[6]:
        for c in range(1, 9):
            ws.cell(i, c).fill = PII

# ------------------------------------------------------ Connect endpoints
rows = [[e["endpoint"], e["path"], e["status"], e["top_level"] if e["status"] == "ok" else "",
         e.get("count") if e.get("count") is not None else "",
         len(e["fields"]), e["answers"], e.get("error", "")[:90]]
        for e in connect["endpoints"]]
sheet("Connect endpoints",
      ["Endpoint", "Path", "Status", "Returns", "Records", "Fields", "What it answers", "Error"],
      rows, [26, 50, 10, 10, 10, 8, 60, 40], wrap_cols=(7,))

os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
wb.save(OUT)
print(f"{OUT}\nConnect {n_connect} fields | Open API {n_open} fields | {len(questions)} answerable questions")
